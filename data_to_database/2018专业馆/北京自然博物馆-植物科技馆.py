# coding=utf-8
import os
from pymongo import MongoClient
from docx import Document

filepath = 'E:/data/专业馆/北京自然博物馆.docx'

client = MongoClient("mongodb://serp:serp123456@127.0.0.1:27017")
db = client.serp

document = Document(filepath)  # 读入文件

table = document.tables[0]
venueName = '北京自然博物馆-植物科技馆'

datas = []
print(venueName, " start reading table")
for i in range(4,6):#从表格第5行开始循环读取表格数据 len(table.rows)
    print(venueName, "'s ", i-3)
    n = 1
    xh = table.cell(i, n-1).text
    zpmc = table.cell(i, n).text
    jbnr = table.cell(i, n+1).text
    ztmc = table.cell(i, n+2).text
    zsjs = table.cell(i, n+3).text  # 数字
    jyx = table.cell(i, n+4).text
    kxx = table.cell(i, n+5).text
    qwx = table.cell(i, n+6).text
    ysx = table.cell(i, n+7).text
    cyx = table.cell(i, n+8).text
    czjd = table.cell(i, n+9).text
    zttc = table.cell(i, n+10).text
    gdmq = table.cell(i, n+11).text
    fgxm = table.cell(i, n+12).text
    wzsmzxgf = table.cell(i, n+13).text     # 数字
    szmt = table.cell(i, n+14).text     # 数字
    zcsjzyxs = table.cell(i, n+15).text     # 数字
    kxzs = table.cell(i, n+16).text     # 数字
    zcnrxz = table.cell(i, n+17).text       # 数字

    zsjs_new = ""
    for index,letter in enumerate(zsjs):
        if letter == '1':
            zsjs_new += "符合主题精神"
        elif letter == '2':
            zsjs_new += "线条优雅、色彩和谐"
        elif letter == '3':
            zsjs_new += "结构牢固"
        elif letter == '4':
            zsjs_new += "避免炫光"
        elif letter == '5':
            zsjs_new += "耐磨、光滑、无缺损、无尖锐棱角"
        elif letter == '6':
            zsjs_new += "操作按键种类、规格统一"
        if index!=(len(zsjs)-1) and letter!='/':
            zsjs_new += "/"

    wzsmzxgf_new = ""
    for index,letter in enumerate(wzsmzxgf):
        if letter == '1':
            wzsmzxgf_new += "下一级文字说明服从于上一级文字说明"
        elif letter == '2':
            wzsmzxgf_new += "字体、文字量、格式等内容上统一"
        elif letter == '3':
            wzsmzxgf_new += "文字说明表述精炼、通俗易懂"
        elif letter == '4':
            wzsmzxgf_new += "使用简化字"
        if index!=(len(wzsmzxgf)-1) and letter!='/':
            wzsmzxgf_new += "/"

    szmt_new = ""
    for index,letter in enumerate(szmt):
        if letter == '1':
            szmt_new += "音频"
        elif letter == '2':
            szmt_new += "视频"
        elif letter == '3':
            szmt_new += "数字媒体触摸屏"
        elif letter == '4':
            szmt_new += "虚拟现实技术"
        elif letter == '5':
            szmt_new += "情景交互数字媒体"
        if index!=(len(szmt)-1) and letter!='/':
            szmt_new += "/"

    zcsjzyxs_new = ""
    for index,letter in enumerate(zcsjzyxs):
        if letter == '1':
            zcsjzyxs_new += "图文板"
        elif letter == '2':
            zcsjzyxs_new += "实物"
        elif letter == '3':
            zcsjzyxs_new += "模型"
        elif letter == '4':
            zcsjzyxs_new += "原复场景或景箱"
        elif letter == '5':
            zcsjzyxs_new += "机电互动展项"
        elif letter == '6':
            zcsjzyxs_new += "多媒体"
        elif letter == '7':
            zcsjzyxs_new += "影院与剧场"
        if index!=(len(zcsjzyxs)-1) and letter!='/':
            zcsjzyxs_new += "/"

    kxzs_new = ""
    for index,letter in enumerate(kxzs):
        if letter == '1':
            kxzs_new += "现象"
        elif letter == '2':
            kxzs_new += "常识"
        elif letter == '3':
            kxzs_new += "高科技"
        elif letter == '4':
            kxzs_new += "观念"
        elif letter == '5':
            kxzs_new += "生产"
        if index!=(len(kxzs)-1) and letter!='/':
            kxzs_new += "/"

    zcnrxz_new = ""
    for index,letter in enumerate(zcnrxz):
        if letter == '1':
            zcnrxz_new += "科学"
        elif letter == '2':
            zcnrxz_new += "技术"
        elif letter == '3':
            zcnrxz_new += "社会"
        elif letter == '4':
            zcnrxz_new += "环境"
        if index!=(len(zcnrxz)-1) and letter!='/':
            zcnrxz_new += "/"

    imgList = []
    videoList = []

    item = {"name": zpmc, "content": jbnr, "room": ztmc, "style_design": zsjs_new,
            "selection_principle*educative_nature": jyx, "selection_principle*science_nature": kxx,
            "selection_principle*interesting_nature": qwx, "selection_principle*artistry_nature": ysx,
            "selection_principle*participative_nature": cyx, "selection_principle*operation_nature": czjd,
            "writing_principle*topic_prominence": zttc, "writing_principle*clear_standpoint": gdmq, "writing_principle*style_distinctive": fgxm, "writing_norm": wzsmzxgf_new,
            "digital_media": szmt_new, "exhibit_style": zcsjzyxs_new, "science_knowledge": kxzs_new, "exhibit_content": zcnrxz_new}
data = {"label": "2018专业馆", "venue": venueName, "item": item, "imgList": [], "videoList": []}
datas.append(data)
db.data.insert(datas)
